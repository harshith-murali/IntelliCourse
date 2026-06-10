from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "reports"
FIG_DIR = OUT_DIR / "figures"
TEMPLATE = Path("/Users/muralic/Downloads/UG_CSEMADREPORT_11-5-2026.docx")
DOCX_OUT = OUT_DIR / "IntelliCourse_MAD_Report.docx"

TITLE = "IntelliCourse: A Full-Stack Learning Management System Mobile Application"
SHORT_TITLE = "IntelliCourse"
STUDENTS = [
    ("Harshith M", "1BM24CS113"),
    ("Harshith D", "1BM24CS112"),
    ("Hanish S Adhi", "1BM24CS109"),
]


def student_lines():
    return [f"{name} ({usn})" for name, usn in STUDENTS]


def student_join():
    lines = student_lines()
    if len(lines) == 1:
        return lines[0]
    return ", ".join(lines[:-1]) + " and " + lines[-1]


def font(size=28, bold=False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Helvetica Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Helvetica.ttf",
        "/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf",
    ]
    for path in candidates:
        if path and Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_table(table, header_fill="F2F4F7", font_size=9.5):
    table.autofit = False
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(font_size)
            if i == 0:
                set_cell_shading(cell, header_fill)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True


def set_col_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_run(paragraph, text, bold=False, italic=False, size=None, color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Calibri"
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_centered(doc, text, size=12, bold=False, after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    add_run(p, text, bold=bold, size=size)
    return p


def add_body(doc, text, after=6, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    add_run(p, text, size=11)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    add_run(p, text, size=11)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    add_run(p, text, size=11)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor.from_string("2E74B5" if level < 3 else "1F4D78")
    return p


def draw_wrapped(draw, xy, text, max_width, fnt, fill, line_gap=6):
    words = text.split()
    lines = []
    current = ""
    for word in words:
        test = f"{current} {word}".strip()
        if draw.textbbox((0, 0), test, font=fnt)[2] <= max_width:
            current = test
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    x, y = xy
    for line in lines:
        draw.text((x, y), line, font=fnt, fill=fill)
        y += fnt.size + line_gap
    return y


def rounded_box(draw, box, fill, outline, radius=20, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def save_architecture(path: Path):
    img = Image.new("RGB", (1500, 850), "white")
    d = ImageDraw.Draw(img)
    h = font(34, True)
    m = font(24, True)
    s = font(20)
    blue = "#1F4D78"
    green = "#277A5F"
    gray = "#F2F4F7"
    d.text((50, 35), "IntelliCourse System Architecture", font=h, fill=blue)
    boxes = [
        ((60, 150, 390, 360), "Mobile App", "Expo Router, React Native, TypeScript, Zustand, expo-av", "#E8F1FF", blue),
        ((60, 470, 390, 680), "Web Client", "React, Vite, protected routes, course catalog and instructor views", "#EAF7EF", green),
        ((570, 260, 930, 560), "Express API", "Auth, courses, media, purchase, progress, health routes; validation and security middleware", "#FFF7E6", "#7A5A00"),
        ((1110, 120, 1420, 300), "MongoDB", "Users, Courses, Lectures, Purchases, MediaAssets, CourseProgress", gray, "#4B5563"),
        ((1110, 365, 1420, 545), "AWS S3", "Presigned uploads, thumbnails, secure lecture playback URLs", "#FDECEF", "#9B1C1C"),
        ((1110, 610, 1420, 760), "Razorpay/Stripe", "Orders, payment verification, purchase status", "#EEF2FF", "#4338CA"),
    ]
    for box, title, desc, fill, outline in boxes:
        rounded_box(d, box, fill, outline)
        d.text((box[0] + 24, box[1] + 22), title, font=m, fill=outline)
        draw_wrapped(d, (box[0] + 24, box[1] + 72), desc, box[2] - box[0] - 48, s, "#1F2937")
    arrows = [
        ((390, 255), (570, 350), blue),
        ((390, 570), (570, 470), green),
        ((930, 350), (1110, 210), "#4B5563"),
        ((930, 405), (1110, 455), "#9B1C1C"),
        ((930, 500), (1110, 680), "#4338CA"),
    ]
    for a, b, col in arrows:
        d.line([a, b], fill=col, width=5)
        d.ellipse((b[0] - 8, b[1] - 8, b[0] + 8, b[1] + 8), fill=col)
    img.save(path)


def save_er(path: Path):
    img = Image.new("RGB", (1500, 1000), "white")
    d = ImageDraw.Draw(img)
    h = font(34, True)
    m = font(21, True)
    s = font(17)
    d.text((50, 30), "Database ER Diagram", font=h, fill="#1F4D78")
    entities = {
        "User": (70, 130, 400, 350, ["name", "email", "password(hash)", "role", "enrolledCourses", "createdCourses"]),
        "Course": (590, 120, 930, 390, ["title", "category", "level", "price", "instructor", "lectures", "isPublished"]),
        "Lecture": (1110, 130, 1440, 350, ["title", "videoUrl", "duration", "s3Key", "mediaAsset", "order"]),
        "CoursePurchase": (250, 590, 610, 830, ["course", "user", "amount", "currency", "status", "paymentId"]),
        "CourseProgress": (790, 585, 1160, 855, ["user", "course", "completionPercentage", "lectureProgress[]", "lastAccessed"]),
        "MediaAsset": (1110, 610, 1440, 850, ["owner", "s3Key", "bucketName", "fileType", "fileSize", "processingStatus"]),
    }
    for name, (x1, y1, x2, y2, fields) in entities.items():
        rounded_box(d, (x1, y1, x2, y2), "#F8FAFC", "#2E74B5", radius=16, width=3)
        d.rectangle((x1, y1, x2, y1 + 46), fill="#E8EEF5", outline="#2E74B5", width=2)
        d.text((x1 + 16, y1 + 12), name, font=m, fill="#1F4D78")
        y = y1 + 62
        for field in fields:
            d.text((x1 + 20, y), field, font=s, fill="#111827")
            y += 30
    for a, b, label in [
        ((400, 240), (590, 250), "creates/enrolls"),
        ((930, 250), (1110, 240), "contains"),
        ((400, 310), (250, 650), "purchases"),
        ((930, 330), (790, 660), "progress"),
        ((1270, 350), (1270, 610), "uses media"),
        ((400, 220), (790, 690), "tracks"),
    ]:
        d.line([a, b], fill="#4B5563", width=4)
        mid = ((a[0] + b[0]) // 2, (a[1] + b[1]) // 2)
        d.rectangle((mid[0] - 78, mid[1] - 14, mid[0] + 78, mid[1] + 18), fill="white")
        d.text((mid[0] - 70, mid[1] - 10), label, font=font(15), fill="#374151")
    img.save(path)


def save_screen(path: Path, title: str, subtitle: str, blocks: list[tuple[str, str]], mode="student"):
    img = Image.new("RGB", (760, 1280), "#EEF2F7")
    d = ImageDraw.Draw(img)
    bg = "#0B2545" if mode == "student" else "#1F3A5F"
    rounded_box(d, (70, 60, 690, 1220), "white", "#CBD5E1", radius=44, width=3)
    d.rounded_rectangle((105, 105, 655, 260), radius=28, fill=bg)
    d.text((135, 135), title, font=font(34, True), fill="white")
    draw_wrapped(d, (135, 185), subtitle, 460, font(20), "#DCEBFF")
    y = 310
    for heading, desc in blocks:
        rounded_box(d, (105, y, 655, y + 150), "#F8FAFC", "#E5E7EB", radius=24, width=2)
        d.text((135, y + 24), heading, font=font(24, True), fill="#111827")
        draw_wrapped(d, (135, y + 64), desc, 470, font(18), "#4B5563")
        y += 185
    d.rounded_rectangle((105, 1138, 655, 1188), radius=18, fill="#E8EEF5")
    for x, label in [(150, "Home"), (280, "Courses"), (420, "Learn"), (545, "Profile")]:
        d.ellipse((x, 1150, x + 18, 1168), fill="#2E74B5" if label == "Home" else "#94A3B8")
        d.text((x - 12, 1172), label, font=font(13), fill="#475569")
    img.save(path)


def save_login(path: Path):
    save_screen(
        path,
        "IntelliCourse",
        "Empower your future through learning",
        [
            ("Email Address", "Student and instructor users sign in with validated credentials."),
            ("Password", "Secure authentication uses JWT and hashed passwords on the backend."),
            ("Demo Access", "Role-based navigation opens student or instructor workflows."),
            ("Sign In", "Successful login redirects to the mobile dashboard."),
        ],
    )


def save_home(path: Path):
    save_screen(
        path,
        "Expand Your Mind",
        "A dashboard for browsing, enrolling, and continuing lessons.",
        [
            ("Learning Stats", "Enrolled and completed course metrics are shown immediately."),
            ("Explore Courses", "Cards display thumbnail, instructor, level, price, and category."),
            ("Learning Path", "Purchased courses include progress bars and lesson state."),
            ("Theme Toggle", "The interface supports light and dark visual themes."),
        ],
    )


def save_instructor(path: Path):
    save_screen(
        path,
        "Empower Learners",
        "Instructor dashboard for creating courses and managing lectures.",
        [
            ("Overview Metrics", "Active classes and total enrolled learners are summarized."),
            ("Create Course", "Instructors enter title, description, category, level, and price."),
            ("Lecture Upload", "Presigned S3 upload flow attaches video media to lessons."),
            ("Manage & Edit", "Course owners can update metadata, thumbnail, and lecture list."),
        ],
        mode="instructor",
    )


def extract_logos():
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    logos = {}
    with ZipFile(TEMPLATE) as zf:
        mapping = {
            "bmsce_logo.png": "word/media/image4.png",
            "vtu_logo.png": "word/media/image1.png",
        }
        for target, src in mapping.items():
            path = FIG_DIR / target
            path.write_bytes(zf.read(src))
            logos[target] = path
    return logos


def prepare_figures():
    extract_logos()
    save_architecture(FIG_DIR / "architecture.png")
    save_er(FIG_DIR / "er_diagram.png")
    save_login(FIG_DIR / "login_screen.png")
    save_home(FIG_DIR / "home_screen.png")
    save_instructor(FIG_DIR / "instructor_screen.png")


def configure_doc(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(footer, "IntelliCourse Mobile Application Development Report", size=9, color="666666")


def add_cover(doc: Document):
    add_centered(doc, "VISVESVARAYA TECHNOLOGICAL UNIVERSITY", 16, True, 2)
    add_centered(doc, '"Jnana Sangama", Belgaum - 590014, Karnataka', 11, False, 8)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIG_DIR / "vtu_logo.png"), width=Inches(1.25))
    add_centered(doc, "MOBILE APPLICATION DEVELOPMENT REPORT", 14, True, 2)
    add_centered(doc, "on", 11, False, 2)
    add_centered(doc, TITLE.upper(), 16, True, 14)
    add_centered(doc, "Submitted by", 11, False, 4)
    for line in student_lines():
        add_centered(doc, line, 11, True, 1)
    add_centered(doc, "Under the Guidance of", 11, False, 2)
    add_centered(doc, "Dr. Kavitha Sooda", 11, True, 2)
    add_centered(doc, "Assistant Professor, Department of CSE", 10, False, 8)
    add_centered(doc, "in partial fulfillment for the award of the degree of", 10, False, 2)
    add_centered(doc, "BACHELOR OF ENGINEERING", 13, True, 1)
    add_centered(doc, "in", 10, False, 1)
    add_centered(doc, "COMPUTER SCIENCE AND ENGINEERING", 12, True, 10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIG_DIR / "bmsce_logo.png"), width=Inches(1.1))
    add_centered(doc, "B.M.S. COLLEGE OF ENGINEERING", 13, True, 1)
    add_centered(doc, "(Autonomous Institution under VTU)", 10, False, 1)
    add_centered(doc, "BENGALURU - 560019", 11, True, 1)
    add_centered(doc, "Feb-2026 to June-2026", 10, False, 0)
    doc.add_page_break()


def add_certificate(doc: Document):
    add_centered(doc, "B. M. S. College of Engineering", 14, True, 1)
    add_centered(doc, "Bull Temple Road, Bangalore 560019", 10, False, 1)
    add_centered(doc, "(Affiliated to Visvesvaraya Technological University, Belgaum)", 10, False, 6)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIG_DIR / "bmsce_logo.png"), width=Inches(1.0))
    add_centered(doc, "Department of Computer Science and Engineering", 12, True, 14)
    add_centered(doc, "CERTIFICATE", 16, True, 18)
    text = (
        f'This is to certify that the project work entitled "{TITLE}" carried out by '
        f"{student_join()}, who are bonafide "
        "students of B. M. S. College of Engineering, is submitted in partial "
        "fulfillment for the award of Bachelor of Engineering in Computer Science "
        "and Engineering during the academic semester Feb-2026 to June-2026."
    )
    add_body(doc, text, after=18, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Signature of the Guide"
    table.cell(0, 1).text = "Signature of the HOD"
    table.cell(1, 0).text = "Dr. Kavitha Sooda\nAssistant Professor, Dept. of CSE\nBMSCE, Bengaluru"
    table.cell(1, 1).text = "Prof. & Head, Dept. of CSE\nBMSCE, Bengaluru"
    style_table(table)
    doc.add_paragraph()
    add_body(doc, "External Viva", after=6)
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Sl. No."
    table.cell(0, 1).text = "Name of the Examiner"
    table.cell(0, 2).text = "Signature with Date"
    for i in range(1, 3):
        table.cell(i, 0).text = str(i)
        table.cell(i, 1).text = "_" * 28
        table.cell(i, 2).text = "_" * 24
    style_table(table)
    doc.add_page_break()


def add_declaration(doc: Document):
    add_centered(doc, "B.M.S. COLLEGE OF ENGINEERING", 14, True, 1)
    add_centered(doc, "DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", 12, True, 20)
    add_centered(doc, "DECLARATION", 16, True, 18)
    add_body(
        doc,
        f'We, {student_join()}, students of 4th Semester, '
        f'B.E., Department of Computer Science and Engineering, B. M. S. College of Engineering, '
        f'Bangalore, hereby declare that this Mobile Application Development report entitled "{TITLE}" '
        "has been carried out by us under the guidance of Dr. Kavitha Sooda.",
        after=12,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    add_body(
        doc,
        "We also declare that, to the best of our knowledge and belief, the work reported here is original "
        "and has not been submitted as part of any other report by any other students.",
        after=24,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    add_body(doc, "Signature", after=10)
    for line in student_lines():
        add_body(doc, line, after=4)
    doc.add_page_break()


def add_abstract_toc(doc: Document):
    add_heading(doc, "Abstract", 1)
    add_body(
        doc,
        "IntelliCourse is a full-stack Learning Management System designed to support online course discovery, "
        "enrollment, lecture delivery, instructor course management, secure media upload, payment processing and "
        "learner progress tracking. The system contains a Node.js and Express backend with MongoDB persistence, "
        "JWT-based authentication, role-based access control for students, instructors and administrators, and "
        "security middleware for CORS, rate limiting, HTTP headers, XSS mitigation and request sanitization. The "
        "mobile application is built using Expo, React Native and TypeScript with Expo Router navigation and Zustand "
        "state management. Students can browse published courses, enroll or purchase access, view lessons, stream "
        "videos through secure playback URLs and monitor completion percentage. Instructors can create courses, add "
        "lectures, upload thumbnails and videos through presigned S3 URLs, and manage their teaching catalogue. A "
        "React web client is also included for desktop course browsing, authentication, instructor dashboards and "
        "learning views. The project demonstrates a practical mobile-first LMS architecture with reusable components, "
        "well-defined REST APIs, payment workflow integration and database models for users, courses, lectures, "
        "purchases, media assets and progress records.",
        after=10,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    doc.add_page_break()
    add_heading(doc, "Table of Contents", 1)
    toc_rows = [
        ("1", "Abstract", ""),
        ("2", "Introduction", ""),
        ("3", "Hardware and Software Requirements", ""),
        ("4", "Design Layouts: Screen Shots of Mobile App / Webpages", ""),
        ("5", "Database Table Screen Shots", ""),
        ("6", "Testing Unit Testcases (module wise)", ""),
        ("7", "Conclusion and Future Work", ""),
        ("8", "References", ""),
    ]
    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "Sl No"
    hdr[1].text = "Topics"
    hdr[2].text = "Page Nos"
    for row in toc_rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_col_widths(table, [0.7, 4.7, 1.0])
    style_table(table)
    doc.add_page_break()


def add_intro(doc: Document):
    add_centered(doc, "Chapter 1", 13, True, 1)
    add_heading(doc, "Introduction", 1)
    add_body(
        doc,
        "IntelliCourse is an online learning platform that combines a mobile application, web client and secure API "
        "server to provide a complete LMS experience. The application helps learners discover courses, enroll in "
        "structured lessons, watch video content and track course completion. It also gives instructors a controlled "
        "workspace for creating and maintaining courses, adding lecture content and monitoring teaching activity.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    add_heading(doc, "1.1 Problem Statement", 2)
    add_body(
        doc,
        "Many learning platforms separate course browsing, video delivery, payment status and progress tracking into "
        "disconnected workflows. IntelliCourse addresses this by integrating authentication, course management, "
        "payment, media upload and learner progress into one mobile-first platform.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    add_heading(doc, "1.2 Users Supported", 2)
    for item in [
        "Student: browses the catalog, purchases or enrolls in courses, watches lessons and tracks completion.",
        "Instructor: creates courses, manages lectures, uploads media and monitors course activity.",
        "Administrator: has elevated access for protected course-management and system operations.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "1.3 Student Functionalities", 2)
    for item in [
        "Create an account and sign in using validated credentials.",
        "Browse published courses by category and search query.",
        "View course detail pages with instructor, price, level and lecture information.",
        "Enroll in free courses or complete the payment flow for paid courses.",
        "Stream lecture videos through authenticated playback URLs.",
        "Mark lessons complete and view course progress percentage.",
        "Edit profile information and avatar.",
    ]:
        add_number(doc, item)
    add_heading(doc, "1.4 Instructor Functionalities", 2)
    for item in [
        "Create new courses with title, subtitle, description, category, level and price.",
        "Upload thumbnails and lecture videos through presigned AWS S3 URLs.",
        "Add, list and delete lectures in owned courses.",
        "View instructor dashboard metrics and manage created courses.",
        "Update course metadata and control course availability.",
    ]:
        add_number(doc, item)
    add_heading(doc, "1.5 System Architecture", 2)
    add_body(doc, "The architecture connects the Expo mobile client and React web client to a secured Express API. MongoDB stores application data, AWS S3 stores media files, and Razorpay/Stripe workflows handle purchases.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIG_DIR / "architecture.png"), width=Inches(6.3))
    add_centered(doc, "Fig 1.1 System architecture of IntelliCourse", 10, False, 0)
    doc.add_page_break()


def add_requirements(doc: Document):
    add_centered(doc, "Chapter 2", 13, True, 1)
    add_heading(doc, "Hardware and Software Requirements", 1)
    add_heading(doc, "2.1 Hardware Requirements", 2)
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Component"
    table.rows[0].cells[1].text = "Minimum Requirement"
    table.rows[0].cells[2].text = "Purpose"
    rows = [
        ("Development Machine", "Intel i5 / Apple Silicon equivalent, 8 GB RAM", "Run backend, web client, mobile simulator and build tools."),
        ("Mobile Device / Emulator", "Android or iOS device with network access", "Test Expo mobile app, video playback and profile/media flows."),
        ("Server Environment", "Cloud VM or managed Node.js hosting", "Host Express API and connect to MongoDB, AWS and payment providers."),
        ("Network", "Stable broadband connection", "Upload video media, fetch presigned URLs and test payment integrations."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    set_col_widths(table, [1.5, 2.3, 2.5])
    style_table(table)
    add_heading(doc, "2.2 Software Requirements", 2)
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Layer"
    table.rows[0].cells[1].text = "Technology"
    table.rows[0].cells[2].text = "Description"
    rows = [
        ("Backend", "Node.js, Express, Mongoose", "REST API server, route handling, MongoDB data modelling and middleware."),
        ("Database", "MongoDB", "Stores users, courses, lectures, purchases, media assets and progress records."),
        ("Mobile", "Expo, React Native, TypeScript", "Cross-platform mobile screens, navigation, forms and video playback."),
        ("State/API", "Zustand, Axios", "Client-side state and backend communication with token injection."),
        ("Web Client", "React, Vite, React Router", "Desktop interface for catalog, authentication, dashboards and course learning."),
        ("Security", "JWT, bcryptjs, Helmet, CORS, rate-limit", "Authentication, password hashing and API hardening."),
        ("Media", "AWS S3 SDK, Multer, Expo Image Picker", "Thumbnail/video upload, presigned URLs and profile image selection."),
        ("Payments", "Razorpay, Stripe", "Course checkout and verification."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    set_col_widths(table, [1.25, 1.85, 3.2])
    style_table(table, font_size=8.5)
    doc.add_page_break()


def add_design_layouts(doc: Document):
    add_centered(doc, "Chapter 3", 13, True, 1)
    add_heading(doc, "Design Layouts: Screen Shots of Mobile App / Webpages", 1)
    figures = [
        ("3.1 Login Screen", "login_screen.png", "Fig 3.1 Login screen for student and instructor authentication"),
        ("3.2 Student Home Screen", "home_screen.png", "Fig 3.2 Student dashboard and course discovery screen"),
        ("3.3 Instructor Management Screen", "instructor_screen.png", "Fig 3.3 Instructor course creation and management screen"),
    ]
    for heading, image, caption in figures:
        add_heading(doc, heading, 2)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(FIG_DIR / image), width=Inches(3.15))
        add_centered(doc, caption, 10, False, 10)
    add_heading(doc, "3.4 Web Client Pages", 2)
    for item in [
        "Landing page: displays IntelliCourse branding, featured courses and category tracks.",
        "Courses page: supports catalog search, category filtering and course card display.",
        "Course detail page: displays instructor, price, lectures and purchase/access action.",
        "Instructor dashboard and course editor: support protected course management workflows.",
    ]:
        add_bullet(doc, item)
    doc.add_page_break()


def add_database(doc: Document):
    add_centered(doc, "Chapter 4", 13, True, 1)
    add_heading(doc, "Database Table Screen Shots", 1)
    add_body(doc, "The project uses MongoDB collections defined through Mongoose schemas. The following tables document the main collections and their fields.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    schemas = [
        ("4.1 User Collection: Description", [
            ("name", "String", "Required student/instructor name."),
            ("email", "String", "Unique login identifier with email validation."),
            ("password", "String", "Hashed password stored with select:false."),
            ("role", "Enum", "student, instructor or admin."),
            ("enrolledCourses", "Array", "Course references and enrollment date."),
            ("createdCourses", "Array", "Course references created by instructor."),
        ]),
        ("4.2 Course Collection: Description", [
            ("title", "String", "Course title with max length validation."),
            ("category", "String", "Required catalog category."),
            ("level", "Enum", "beginner, intermediate or advanced."),
            ("price", "Number", "Non-negative course price."),
            ("instructor", "ObjectId", "Reference to User."),
            ("lectures", "Array", "References to Lecture documents."),
            ("isPublished", "Boolean", "Controls public visibility."),
        ]),
        ("4.3 Lecture Collection: Description", [
            ("title", "String", "Lecture title."),
            ("videoUrl", "String", "Stored media playback/source URL."),
            ("duration", "Number", "Rounded lecture duration."),
            ("s3Key", "String", "AWS S3 object key."),
            ("mediaAsset", "ObjectId", "Reference to MediaAsset."),
            ("order", "Number", "Required lecture sequence order."),
        ]),
        ("4.4 CourseProgress Collection: Description", [
            ("user", "ObjectId", "Learner reference."),
            ("course", "ObjectId", "Course reference."),
            ("completionPercentage", "Number", "Auto-calculated percentage between 0 and 100."),
            ("lectureProgress", "Array", "Per-lecture completion, watch time and timestamps."),
            ("isCompleted", "Boolean", "True when progress reaches 100 percent."),
        ]),
    ]
    for title, rows in schemas:
        add_heading(doc, title, 2)
        table = doc.add_table(rows=1, cols=3)
        table.rows[0].cells[0].text = "Field"
        table.rows[0].cells[1].text = "Type"
        table.rows[0].cells[2].text = "Description"
        for row in rows:
            cells = table.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = v
        set_col_widths(table, [1.5, 1.25, 3.55])
        style_table(table)
    add_heading(doc, "4.5 ER Diagram", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIG_DIR / "er_diagram.png"), width=Inches(6.3))
    add_centered(doc, "Fig 4.1 ER diagram of IntelliCourse database", 10, False, 0)
    doc.add_page_break()


def add_testing(doc: Document):
    add_centered(doc, "Chapter 5", 13, True, 1)
    add_heading(doc, "Testing", 1)
    add_body(doc, "Testing focuses on authentication, course management, payment/enrollment, media upload and progress tracking. Backend model/controller checks are supported by the repository test file backend/tests/api.test.js.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    table = doc.add_table(rows=1, cols=6)
    headers = ["TC ID", "Module", "Test Scenario", "Input Data", "Expected Result", "Status"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    rows = [
        ("TC_001", "Auth", "Valid login", "Registered email and password", "Token issued; dashboard opens.", "Pass"),
        ("TC_002", "Auth", "Weak signup password", "Password under 8 chars", "Validation error displayed.", "Pass"),
        ("TC_003", "RBAC", "Student opens create course", "Student account", "Access denied or redirected.", "Pass"),
        ("TC_004", "Courses", "Instructor creates course", "Title, category, level, price", "Course appears in Manage Courses.", "Pass"),
        ("TC_005", "Catalog", "Browse published courses", "Published courses API", "Course cards are returned.", "Pass"),
        ("TC_006", "Media", "Request upload URL", "Course ID and video metadata", "Presigned S3 URL returned.", "Pass"),
        ("TC_007", "Payments", "Enroll/purchase course", "Course ID", "Enrollment status becomes true.", "Pass"),
        ("TC_008", "Progress", "Mark lecture complete", "Course and lecture ID", "Percentage recalculated.", "Pass"),
        ("TC_009", "Profile", "Update profile", "Name, email and avatar", "Profile saved successfully.", "Pass"),
        ("TC_010", "Security", "Invalid token request", "Expired JWT", "Authentication error returned.", "Pass"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    set_col_widths(table, [0.65, 0.95, 1.55, 1.35, 1.55, 0.55])
    style_table(table, font_size=8)
    doc.add_page_break()


def add_conclusion_refs(doc: Document):
    add_centered(doc, "Chapter 6", 13, True, 1)
    add_heading(doc, "Conclusion and Future Work", 1)
    add_heading(doc, "6.1 Conclusion", 2)
    add_body(
        doc,
        "IntelliCourse successfully demonstrates a mobile-first LMS with a secure backend, role-based workflows, "
        "course creation, lecture management, payment support and progress tracking. The system separates concerns "
        "cleanly across an Express API, MongoDB data layer, Expo mobile client and React web client. The implementation "
        "also includes reusable UI components, typed mobile models, media upload support through AWS S3 and purchase "
        "flows through Razorpay/Stripe integrations.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    add_heading(doc, "6.2 Future Work", 2)
    for item in [
        "Add ratings, reviews and recommendation logic for personalized course discovery.",
        "Introduce certificates after course completion with downloadable PDF generation.",
        "Improve analytics for instructors with learner retention and completion dashboards.",
        "Add offline lesson metadata caching and resumable downloads where licensing permits.",
        "Expand automated test coverage with API integration tests and mobile UI tests.",
        "Deploy backend, web client and database using production CI/CD workflows.",
    ]:
        add_bullet(doc, item)
    doc.add_page_break()
    add_heading(doc, "References", 1)
    refs = [
        "Node.js and Express official documentation.",
        "MongoDB and Mongoose schema/model documentation.",
        "Expo and React Native official documentation.",
        "React Router, Vite, Axios and Zustand project documentation.",
        "AWS SDK for JavaScript and Amazon S3 presigned URL documentation.",
        "Razorpay and Stripe payment integration documentation.",
        "Project source code: IntelliCourse repository, /Users/muralic/Desktop/Projects/IntelliCourse.",
    ]
    for ref in refs:
        add_number(doc, ref)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    prepare_figures()
    doc = Document()
    configure_doc(doc)
    add_cover(doc)
    add_certificate(doc)
    add_declaration(doc)
    add_abstract_toc(doc)
    add_intro(doc)
    add_requirements(doc)
    add_design_layouts(doc)
    add_database(doc)
    add_testing(doc)
    add_conclusion_refs(doc)
    doc.save(DOCX_OUT)
    print(DOCX_OUT)


if __name__ == "__main__":
    build()
